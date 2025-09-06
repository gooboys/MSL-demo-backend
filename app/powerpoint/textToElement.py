from pptx.util import Pt

def set_text(shape, text, font_name="DM Sans", font_size=10, bold=False):
  tf = shape.text_frame
  tf.clear()
  run = tf.paragraphs[0].add_run()
  run.text = text
  font = run.font
  font.name = font_name
  font.size = Pt(font_size)
  font.bold = bold